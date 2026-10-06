from __future__ import annotations

import csv
import io
import json
import zipfile
from typing import Literal
from uuid import UUID, uuid4

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import __version__
from app.persistence.models import ComponentModel, FunctionNodeModel, ProjectModel, RelationModel, RevisionModel, SequenceAllocationModel
from app.rds.allocator import allocation_scope
from app.services.project_service import ProjectService
from app.standards.loader import RulesetRegistry

FLOC_HEADERS = ["RDS", "Component", "Function", "Class", "Parent", "Description", "Status", "Warnings"]


def floc_rows(service: ProjectService, project_id: UUID) -> list[list[str]]:
    project = service.get_project(project_id)  # compiles current state immediately
    rows = [
        [
            item["designation"],
            item["label"],
            f"M{item['function_number']}" if item["function_number"] is not None else "",
            item["class_code"],
            item["parent_label"] or "",
            item["description"],
            "Valid" if item["valid"] else "Invalid",
            "; ".join(issue["message"] for issue in item["issues"]),
        ]
        for item in project["components"]
    ]
    return sorted(rows, key=lambda row: row[0])


def engineering_export(service: ProjectService, project_id: UUID, format: Literal["json", "csv", "xlsx", "pdf"]) -> tuple[bytes, str]:
    project = service.get_project(project_id)
    rows = floc_rows(service, project_id)
    if format == "json":
        return json.dumps({"project": project["name"], "floc": [dict(zip(FLOC_HEADERS, row, strict=True)) for row in rows]}, default=str, indent=2).encode(), "application/json"
    if format == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(FLOC_HEADERS)
        writer.writerows(rows)
        return stream.getvalue().encode("utf-8-sig"), "text/csv"
    if format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "FLoC"
        sheet.append(FLOC_HEADERS)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="5B3FD6")
        widths = [24, 24, 12, 10, 24, 42, 12, 42]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        output = io.BytesIO()
        book.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    data = [[Paragraph(header, styles["BodyText"]) for header in FLOC_HEADERS]]
    data += [[Paragraph(str(cell), styles["BodyText"]) for cell in row] for row in rows]
    table = Table(data, repeatRows=1, colWidths=[31 * mm, 30 * mm, 18 * mm, 14 * mm, 27 * mm, 47 * mm, 18 * mm, 62 * mm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5B3FD6")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8B4C4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F0F5")]),
            ]
        )
    )
    document.build([Paragraph(f"AutoRDS — FLoC — {project['name']}", styles["Title"]), Spacer(1, 5 * mm), table])
    return output.getvalue(), "application/pdf"


def export_portable(service: ProjectService, project_id: UUID) -> bytes:
    project = service.get_project(project_id)
    ruleset = RulesetRegistry.get(project["ruleset_id"])
    history = service.session.scalars(select(RevisionModel).where(RevisionModel.project_id == str(project_id)).order_by(RevisionModel.timestamp)).all()
    manifest = {"format": "AutoRDS Project", "format_version": 1, "application_version": __version__, "ruleset": ruleset.id, "ruleset_version": ruleset.version}
    files = {
        "manifest.json": manifest,
        "project.json": {key: project[key] for key in ("id", "name", "description", "ruleset_id")},
        "functions.json": project["functions"],
        "components.json": project["components"],
        "relations.json": project["relations"],
        "history.json": [{column.name: getattr(item, column.name) for column in item.__table__.columns} for item in history],
        "ruleset.json": ruleset.model_dump(),
        "queries.json": [],
    }
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in files.items():
            archive.writestr(name, json.dumps(payload, default=str, indent=2))
    return output.getvalue()


def import_portable(session: Session, content: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        names = set(archive.namelist())
        required = {"manifest.json", "project.json", "functions.json", "components.json", "relations.json", "history.json", "ruleset.json", "queries.json"}
        if not required.issubset(names):
            raise ValueError("Portable project is missing required files")

        def load(name: str):
            return json.loads(archive.read(name))

        manifest, project, functions, components, relations = load("manifest.json"), load("project.json"), load("functions.json"), load("components.json"), load("relations.json")
    if manifest.get("format") != "AutoRDS Project" or manifest.get("format_version") != 1:
        raise ValueError("Unsupported AutoRDS project format")
    RulesetRegistry.get(project["ruleset_id"])
    new_project_id = str(uuid4())
    id_map: dict[str, str] = {}
    session.add(ProjectModel(id=new_project_id, name=project["name"], description=project.get("description", ""), ruleset_id=project["ruleset_id"]))
    session.flush()
    for item in functions:
        id_map[str(item["id"])] = str(uuid4())
    for item in components:
        id_map[str(item["id"])] = str(uuid4())
    for item in functions:
        session.add(FunctionNodeModel(id=id_map[str(item["id"])], project_id=new_project_id, number=item["number"], label=item["label"], description=item.get("description", ""), parent_id=id_map.get(str(item.get("parent_id")))))
    session.flush()
    for item in components:
        model = ComponentModel(
            id=id_map[str(item["id"])],
            project_id=new_project_id,
            function_id=id_map.get(str(item["function_id"])),
            parent_component_id=id_map.get(str(item.get("parent_component_id"))),
            class_code=item["class_code"],
            allocated_number=item["allocated_number"],
            label=item["label"],
            description=item.get("description", ""),
            source_type=item.get("source_type", "import"),
            source_query_id=None,
        )
        session.add(model)
        session.flush()
        domain = ProjectService(session)._component_domain(model)
        session.add(SequenceAllocationModel(scope_id=allocation_scope(domain), class_code=model.class_code, component_id=model.id, assigned_number=model.allocated_number, active=True))
    for item in relations:
        session.add(RelationModel(id=str(uuid4()), project_id=new_project_id, source_component_id=id_map[str(item["source_component_id"])], target_component_id=id_map[str(item["target_component_id"])], relation_type=item["relation_type"]))
    session.commit()
    service = ProjectService(session)
    service._validate_and_commit(UUID(new_project_id))
    return service.get_project(UUID(new_project_id))
